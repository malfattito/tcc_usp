"""Preenche o Template Projeto de Pesquisa (251, 252).docx com o conteúdo do projeto,
preservando a formatação do template (Arial 11, justificado, recuo e espaçamento do modelo).

Uso: python3 scripts/gerar_projeto_pesquisa.py
Saída: projeto_pesquisa/Projeto_de_Pesquisa_Silvano_Malfatti.docx
"""
import copy, os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(RAIZ, "Template Projeto de Pesquisa (251, 252).docx")
SAIDA = os.path.join(RAIZ, "projeto_pesquisa", "Projeto_de_Pesquisa_Silvano_Malfatti.docx")

ALUNO = "Silvano Maneck Malfatti"
ORIENTADOR = "[nome do(a) orientador(a) designado(a)]"
CURSO = "MBA em Data Science e Analytics"
TITULO = "Simulador web de risco fiscal de municípios brasileiros com aprendizado de máquina"

INTRODUCAO = [
    "Os municípios brasileiros respondem por parcela expressiva dos serviços públicos de saúde, educação básica, "
    "assistência social e infraestrutura urbana, mas contam com base tributária própria limitada e dependem, em sua "
    "maioria, de transferências da União e dos estados. A Lei Complementar nº 101, de 2000, conhecida como Lei de "
    "Responsabilidade Fiscal [LRF], buscou disciplinar essa realidade ao fixar limites para a despesa com pessoal, "
    "equivalentes a 54% da receita corrente líquida [RCL] no Poder Executivo municipal, e para o endividamento, além "
    "de exigir a publicação periódica de relatórios fiscais (Brasil, 2000). Esses relatórios são hoje consolidados pela "
    "Secretaria do Tesouro Nacional [STN] no Sistema de Informações Contábeis e Fiscais do Setor Público Brasileiro "
    "[SICONFI], que reúne, para os 5.570 municípios, as declarações de contas anuais e os relatórios de gestão fiscal.",

    "Apesar desse arcabouço, a situação fiscal dos municípios permanece heterogênea e, em muitos casos, frágil. "
    "Marconato et al. (2021), ao analisarem dados em painel de 2002 a 2016, concluíram que despesas elevadas com "
    "pessoal deterioram o resultado fiscal e que os piores desempenhos se concentram em municípios pequenos de perfil "
    "industrial. Salomão Neto e Santos (2023), com modelos logit e probit aplicados aos 5.570 municípios entre 1994 e "
    "2017, mostraram que a dependência de transferências correntes eleva a probabilidade de a despesa com pessoal "
    "ultrapassar o limite prudencial da LRF, sobretudo nos municípios menos populosos. Oliveira e Oliveira (2023), ao "
    "aplicarem a todos os municípios a métrica de sustentabilidade proposta na PEC 188/2019, verificaram que mais de "
    "85% deles seriam classificados como insustentáveis. No nível estadual, Lima et al. (2024) associaram as "
    "transferências federais e a receita tributária ao descumprimento dos limites de pessoal entre 2017 e 2022.",

    "Para subsidiar a concessão de garantias da União em operações de crédito, a STN calcula anualmente a Capacidade "
    "de Pagamento [CAPAG] de cada município, que combina indicadores de endividamento, poupança corrente e liquidez em "
    "uma nota que varia de A a D, com metodologia atualizada pela Portaria Normativa MF nº 1.583, de 2023 (Brasil, "
    "2023). Antunes e Silva (2025) adotaram a nota CAPAG como medida da condição financeira de 5.570 municípios entre "
    "2018 e 2021 e identificaram indicadores socioeconômicos associados a ela. Batista et al. (2024) construíram um "
    "índice de condição financeira a partir de vinte indicadores do SICONFI e observaram que a pandemia de Covid-19 "
    "não alterou substancialmente a posição relativa dos municípios de maior porte. Esses estudos têm caráter "
    "descritivo ou explicativo e, em geral, examinam a condição fiscal de forma contemporânea, sem o propósito de "
    "antecipá-la.",

    "Na literatura internacional, técnicas de aprendizado de máquina (\"machine learning\") vêm sendo aplicadas "
    "justamente com esse objetivo preditivo. Gorina et al. (2018) modelaram o estresse fiscal de governos locais "
    "norte-americanos e apontaram as reservas financeiras, a dívida e a composição da receita como os principais "
    "preditores. Antulov-Fantulin et al. (2021), com dados de 7.795 municípios italianos entre 2009 e 2016, obtiveram "
    "previsões fora da amostra com alta taxa de verdadeiros positivos, mesmo diante da raridade dos casos de "
    "insolvência, e constataram que características não financeiras figuram entre os preditores mais relevantes. Liu "
    "et al. (2025) compararam seis algoritmos e diferentes estratégias de balanceamento de classes para governos locais "
    "de 49 estados norte-americanos, com vantagem para métodos de agregação balanceada. No Brasil, Garruti et al. "
    "(2025) aplicaram o algoritmo Random Forest à previsão de estresse fiscal nos municípios de Minas Gerais entre 2016 "
    "e 2020, com acurácia média de 68,2% e predominância das variáveis de liquidez entre as mais importantes, o que "
    "evidencia tanto o potencial quanto as limitações da abordagem quando restrita a um único estado.",

    "Não foram identificados estudos que comparem algoritmos supervisionados para prever, com um ano de "
    "antecedência, a classificação de risco fiscal do conjunto dos municípios brasileiros, com validação temporal e "
    "tratamento explícito do desbalanceamento entre as classes, nem que incorporem o modelo resultante em um simulador "
    "no qual o gestor informe as despesas previstas para o exercício corrente e receba a estimativa de risco antes de "
    "executá-las. Antecipar a deterioração fiscal interessa a gestores municipais, que podem testar decisões de gasto "
    "antes de atingir os limites legais, aos tribunais de contas e à STN, que podem priorizar o acompanhamento, e aos "
    "cidadãos, que ganham um instrumento de controle social. A disponibilidade pública e padronizada dos dados do "
    "SICONFI e da CAPAG, aliada às técnicas de aprendizado supervisionado estudadas no curso, torna o problema "
    "tratável e justifica a realização desta pesquisa.",
]

OBJETIVO = [
    "O objetivo geral é desenvolver um simulador web, baseado em modelos supervisionados de aprendizado de máquina, "
    "que estime com um ano de antecedência o risco fiscal de municípios brasileiros, medido pela nota CAPAG da STN, a "
    "partir dos dados publicados no SICONFI e das despesas previstas informadas pelo gestor.",
    "Como objetivos específicos, pretende-se: (i) consolidar uma base município-ano com variáveis fiscais do SICONFI "
    "e as notas CAPAG de 2019 a 2025; (ii) treinar e comparar, com validação temporal, modelos de regressão "
    "logística, árvore de decisão, Random Forest, \"gradient boosting\" e rede neural artificial; (iii) selecionar o "
    "modelo de melhor desempenho, impor restrições de monotonicidade às variáveis de despesa e interpretar seus "
    "efeitos; (iv) implementar o simulador como aplicação web em Python, com o modelo embarcado e explicação de cada "
    "estimativa; e (v) avaliar o simulador de forma retrospectiva, verificando se teria alertado os municípios que "
    "foram rebaixados para as notas C ou D nos anos de teste.",
]

METODOLOGIA = [
    "A pesquisa será de natureza aplicada e abordagem quantitativa, com objetivos descritivos e explicativos, pois "
    "pretende descrever a situação fiscal dos municípios e identificar as variáveis associadas ao risco futuro. Quanto "
    "ao delineamento, enquadra-se como Implementação de Algoritmos de Machine Learning, modalidade prevista para os "
    "TCCs do MBA em Data Science e Analytics, e utilizará como técnica de obtenção de dados o levantamento de dados "
    "secundários em bases públicas (Gil, 2017).",

    "Os dados serão obtidos de duas fontes da STN. A primeira é a API do SICONFI, da qual serão extraídas, para cada "
    "município e exercício de 2019 a 2024, a Declaração de Contas Anuais [DCA], com receitas e despesas orçamentárias "
    "por natureza e por função, e o Relatório de Gestão Fiscal [RGF] do Poder Executivo referente ao terceiro "
    "quadrimestre, com a despesa total com pessoal, a RCL, a dívida consolidada e a disponibilidade de caixa. A segunda "
    "é o conjunto de dados Capacidade de Pagamento de Municípios, publicado no portal Tesouro Transparente, com as "
    "notas CAPAG dos anos-base 2020 a 2025 e os indicadores que as compõem (Secretaria do Tesouro Nacional [STN], "
    "2026). A população de cada município será obtida do cadastro de entes do próprio SICONFI. A coleta será "
    "automatizada em Python, com registro das datas de acesso e armazenamento dos arquivos brutos para garantir a "
    "reprodutibilidade.",

    "A unidade de análise será o par município-ano. As variáveis explicativas serão expressas como razões, de modo a "
    "neutralizar diferenças de escala entre municípios, entre elas: despesa com pessoal sobre a RCL; receita "
    "tributária própria sobre a receita corrente; transferências correntes sobre a receita corrente; dívida "
    "consolidada sobre a RCL; disponibilidade de caixa líquida sobre a despesa corrente; restos a pagar sobre a "
    "receita corrente; poupança corrente, dada pela diferença entre receita e despesa correntes sobre a receita "
    "corrente; investimentos sobre a despesa total; além de população, região e porte. Serão incluídas defasagens de "
    "um e dois anos das principais razões, para capturar tendência. A variável dependente principal será a nota CAPAG "
    "do ano-base seguinte, dicotomizada em alto risco (notas C e D) e baixo risco (notas A, A+, B e B+). Como variável "
    "dependente alternativa, será testado o indicador de a despesa total com pessoal atingir ou superar o limite "
    "prudencial da LRF no exercício seguinte. Municípios sem nota atribuída pela STN por ausência de entrega das "
    "declarações serão excluídos do rótulo daquele ano, e a mudança metodológica da CAPAG a partir do ano-base 2023 "
    "será controlada por meio de variável indicadora e de análises separadas por período.",

    "No tratamento dos dados, valores ausentes serão imputados pela mediana da unidade da federação quando "
    "inferiores a 5% das observações da variável e, acima disso, a variável será descartada; valores extremos serão "
    "limitados aos percentis 1 e 99; e as variáveis serão padronizadas para os modelos sensíveis à escala. Os "
    "algoritmos avaliados serão a regressão logística, adotada como referência, a árvore de decisão, o Random Forest "
    "(Breiman, 2001), o \"gradient boosting\" e a rede neural artificial do tipo perceptron de múltiplas camadas, "
    "conforme Fávero e Belfiore (2017). Não serão empregadas arquiteturas de aprendizado profundo, pois, em dados "
    "tabulares de dimensão moderada, modelos baseados em árvores apresentam desempenho igual ou superior ao de redes "
    "neurais profundas com menor custo de ajuste (Shwartz-Ziv e Armon, 2022; Grinsztajn et al., 2022), e a rede "
    "neural rasa incluída na comparação permitirá verificar essa evidência no problema estudado. A validação será "
    "temporal, para evitar vazamento de informação futura: os "
    "modelos serão treinados com as observações de 2019 a 2021, cujos rótulos correspondem aos anos-base 2020 a 2022, "
    "ajustados em 2022 com rótulo de 2023 e testados em 2023 e 2024 com rótulos de 2024 e 2025. Os hiperparâmetros "
    "serão escolhidos por validação cruzada estratificada dentro do conjunto de treino. Ao modelo selecionado serão "
    "impostas restrições de monotonicidade, de modo que a probabilidade estimada de alto risco nunca diminua quando "
    "aumentam a despesa com pessoal, a despesa corrente ou a dívida, garantindo respostas coerentes do simulador; o "
    "custo dessa restrição será medido pela diferença de desempenho em relação ao modelo irrestrito. O desbalanceamento entre as "
    "classes será tratado com ponderação inversa à frequência e, de forma comparativa, com subamostragem da classe "
    "majoritária.",

    "O desempenho será medido pela área sob a curva ROC [AUC], pela área sob a curva de precisão e sensibilidade, pela "
    "sensibilidade, pela especificidade, pela acurácia balanceada e pelo escore F1, com intervalos de confiança "
    "obtidos por reamostragem. A interpretação dos modelos será feita por importância de variáveis por permutação e "
    "por valores SHAP, que permitirão identificar, para cada município, os fatores que mais contribuíram para a "
    "previsão. As análises serão realizadas em Python, com as bibliotecas pandas, scikit-learn, lightgbm e shap, e o "
    "código será mantido em repositório público.",

    "O simulador será uma aplicação web desenvolvida em Python, com o framework Streamlit, hospedada em ambiente com "
    "suporte à linguagem, e carregará o modelo selecionado já treinado. O usuário escolherá o município e verá suas "
    "contas mais recentes publicadas no SICONFI; poderá então alterar as rubricas de despesa sob controle da gestão, "
    "a saber, despesa com pessoal, outras despesas correntes, investimentos e novas operações de crédito, enquanto as "
    "demais variáveis, como receitas e transferências, permanecerão fixas no último valor observado ou em projeção "
    "informada pelo usuário. A cada alteração, o simulador apresentará dois resultados: os indicadores de "
    "endividamento, poupança corrente e liquidez calculados de forma determinística segundo a metodologia da CAPAG "
    "para o exercício simulado, e a probabilidade estimada pelo modelo de o município ser classificado como alto "
    "risco no ano-base seguinte, acompanhada dos fatores que mais contribuíram para a estimativa, obtidos dos valores "
    "SHAP. A interface deixará explícito que a estimativa decorre de associações observadas nos dados históricos e "
    "constitui indício, e não garantia, do efeito de uma decisão de gasto.",

    "A avaliação do simulador será retrospectiva e exclusivamente técnica, sem participação de gestores ou de outros "
    "sujeitos de pesquisa. Para cada município que foi rebaixado para as notas C ou D nos anos-base de teste, 2024 e "
    "2025, o simulador será alimentado com as contas do exercício anterior, tal como estavam publicadas à época, e "
    "será verificado se teria emitido alerta de alto risco. Serão calculadas a proporção de rebaixamentos antecipados, "
    "a proporção de alertas emitidos para municípios que mantiveram ou melhoraram a nota e o número de municípios "
    "nos quais uma redução simulada de cinco pontos percentuais na despesa com pessoal teria levado a estimativa "
    "para abaixo do limiar de alerta. Será verificada ainda a coerência das respostas, confirmando que nenhum aumento "
    "simulado de despesa reduz a probabilidade estimada de risco.",

    "Quanto aos aspectos éticos, a pesquisa utilizará exclusivamente informações de acesso público e de domínio "
    "público, referentes a entes federativos e sem dados pessoais, o que dispensa a submissão ao Comitê de Ética em "
    "Pesquisa, conforme os incisos II e III do parágrafo único do artigo 1º da Resolução CNS nº 510, de 2016 (Brasil, "
    "2016). Nenhuma etapa, inclusive a avaliação do simulador, envolverá coleta de dados com pessoas. O Formulário de "
    "Direcionamento Ético será preenchido e validado com o orientador. Os resultados serão "
    "apresentados como estimativas de risco e não como julgamento sobre a gestão de prefeitos ou servidores, e a "
    "ausência de nota por não entrega de declarações será sempre distinguida de uma nota desfavorável. Durante a "
    "preparação deste projeto, o autor utilizou o assistente de inteligência artificial Claude, da Anthropic, para "
    "organizar as referências, estruturar o código de coleta de dados e revisar o texto. Após utilizar essa ferramenta, "
    "o autor revisou e editou o conteúdo conforme necessário e assume total responsabilidade pelo conteúdo deste "
    "projeto e do trabalho dele decorrente.",
]

RESULTADOS_ESPERADOS = [
    "Espera-se que ao menos um dos modelos baseados em árvores alcance AUC igual ou superior a 0,80 e sensibilidade "
    "igual ou superior a 0,75 para a classe de alto risco no conjunto de teste, superando a regressão logística em "
    "pelo menos cinco pontos percentuais de AUC, resultado compatível com o desempenho relatado para municípios "
    "italianos e superior ao obtido com Random Forest em Minas Gerais. Espera-se também que a ponderação de classes "
    "melhore a sensibilidade sem perda relevante de especificidade.",
    "Na análise de importância, espera-se que a dependência de transferências correntes, a razão entre despesa com "
    "pessoal e RCL, a disponibilidade de caixa líquida e o porte populacional figurem entre as variáveis mais "
    "relevantes, com efeito positivo das duas primeiras sobre a probabilidade de alto risco e negativo das duas "
    "últimas. Espera-se ainda forte persistência da nota de um ano para o outro, de modo que as defasagens das razões "
    "fiscais ganhem peso na previsão, e que o desempenho seja menor para os municípios de até 10 mil habitantes, "
    "nos quais a volatilidade das receitas é maior.",
    "Espera-se que as restrições de monotonicidade reduzam a AUC do modelo selecionado em no máximo 0,02, "
    "preservando a coerência das respostas do simulador. Na avaliação retrospectiva, espera-se que o simulador teria "
    "alertado ao menos 70% dos municípios rebaixados para C ou D nos anos-base 2024 e 2025, com alertas emitidos para "
    "no máximo 30% dos municípios que mantiveram ou melhoraram a nota, e que responda a cada simulação em menos de "
    "um segundo.",
    "Como produto, espera-se entregar o simulador em funcionamento na web, com a situação fiscal, a estimativa de "
    "risco e a simulação de cenários de despesa para cada um dos 5.570 municípios, além de uma base consolidada e um "
    "código reprodutíveis, que possam ser atualizados a cada nova divulgação da CAPAG.",
]

# Cronograma: 10 meses (colunas) x 12 atividades (linhas). Ajustar ao calendário oficial da turma.
MESES = ["out/26", "nov/26", "dez/26", "jan/27", "fev/27", "mar/27", "abr/27", "mai/27", "jun/27", "jul/27"]
ATIVIDADES = [
    ("Revisão da literatura e redação do Projeto de Pesquisa", [1, 2]),
    ("Coleta automatizada dos dados do SICONFI e da CAPAG", [1, 2]),
    ("Tratamento, consolidação da base e construção das variáveis", [2, 3]),
    ("Análise exploratória e definição dos rótulos", [3, 4]),
    ("Treinamento e validação temporal dos modelos", [4, 5]),
    ("Redação e entrega dos Resultados Preliminares", [5]),
    ("Ajuste de hiperparâmetros, monotonicidade e interpretação (SHAP)", [6]),
    ("Desenvolvimento do simulador web e avaliação retrospectiva", [6, 7]),
    ("Redação do TCC: resultados, discussão e conclusão", [7, 8]),
    ("Revisão pelo orientador e entrega do TCC", [8]),
    ("Preparação dos slides e defesa", [9]),
    ("Entrega da versão revisada do TCC", [10]),
]

REFERENCIAS = [
    "Antulov-Fantulin, N.; Lagravinese, R.; Resce, G. 2021. Predicting bankruptcy of local government: a machine "
    "learning approach. Journal of Economic Behavior and Organization 183: 681-699.",
    "Antunes, R.M.C.; Silva, C.A.T. 2025. Indicadores socioeconômicos determinantes da condição financeira nos "
    "municípios brasileiros. Desenvolvimento em Questão 23(62): e15733.",
    "Batista, A.P.; Carvalho, T.F.M.; Macedo, M.A.S.; Cruz, C.F. 2024. Condição financeira dos municípios brasileiros "
    "no contexto da pandemia da Covid-19. Administração Pública e Gestão Social 16(1). Disponível em: "
    "<https://www.redalyc.org/articulo.oa?id=351576210002>. Acesso em: 16 out. 2026.",
    "Brasil. 2000. Lei Complementar nº 101, de 4 de maio de 2000. Estabelece normas de finanças públicas voltadas "
    "para a responsabilidade na gestão fiscal e dá outras providências. Diário Oficial da União, Brasília, 5 maio "
    "2000. Seção 1.",
    "Brasil. 2016. Conselho Nacional de Saúde. Resolução nº 510, de 7 de abril de 2016. Dispõe sobre as normas "
    "aplicáveis a pesquisas em Ciências Humanas e Sociais. Diário Oficial da União, Brasília, 24 maio 2016. Seção 1.",
    "Brasil. 2023. Ministério da Fazenda. Portaria Normativa MF nº 1.583, de 22 de dezembro de 2023. Dispõe sobre a "
    "metodologia de cálculo da capacidade de pagamento de Estados, Distrito Federal e Municípios. Diário Oficial da "
    "União, Brasília, dez. 2023. Seção 1.",
    "Breiman, L. 2001. Random forests. Machine Learning 45(1): 5-32.",
    "Fávero, L.P.; Belfiore, P. 2017. Manual de Análise de Dados. Elsevier, Rio de Janeiro, RJ, Brasil.",
    "Garruti, D.V.T.; Barboza, F.L.M.; Diniz, J.A. 2025. Utilização de inteligência artificial na previsão do estresse "
    "fiscal: evidências para municípios brasileiros. Advances in Scientific and Applied Accounting 17(3): 152-166.",
    "Gil, A.C. 2017. Como Elaborar Projetos de Pesquisa. 6ed. Atlas, São Paulo, SP, Brasil.",
    "Gorina, E.; Maher, C.; Joffe, M. 2018. Local fiscal distress: measurement and prediction. Public Budgeting & "
    "Finance 38(1): 72-94.",
    "Grinsztajn, L.; Oyallon, E.; Varoquaux, G. 2022. Why do tree-based models still outperform deep learning on "
    "typical tabular data? Advances in Neural Information Processing Systems 35: 507-520.",
    "Lima, C.A.M.C.; Boente, D.R.; Stefanelli, N.O. 2024. Determinantes do cumprimento das despesas com pessoal e do "
    "resultado primário nos estados brasileiros: uma análise de 2017 a 2022. Revista de Administração Pública 58(3): "
    "e2023-0321.",
    "Liu, R.; Li, H.; Yoon, K. 2025. Using machine learning algorithms to improve fiscal distress prediction models: "
    "the case of U.S. local governments. Journal of Information Systems 39(3): 131-155.",
    "Marconato, M.; Parré, J.L.; Coelho, M.H. 2021. Dinâmica financeira dos municípios brasileiros. Revista de "
    "Administração Pública 55(2): 378-394.",
    "Oliveira, D.T.; Oliveira, A.G. 2023. (In)sustentabilidade financeira municipal: a frágil metodologia proposta "
    "pela PEC do Pacto Federativo. Revista de Administração Pública 57(5): e2023-0012.",
    "Salomão Neto, B.A.; Santos, J.F.C. 2023. Dependência fiscal, ilusão fiscal e dificuldades orçamentárias nos "
    "municípios brasileiros: diagnóstico via modelos logit/probit. Planejamento e Políticas Públicas 64. Disponível "
    "em: <https://doi.org/10.38116/ppp64art4>. Acesso em: 16 out. 2026.",
    "Shwartz-Ziv, R.; Armon, A. 2022. Tabular data: deep learning is not all you need. Information Fusion 81: 84-90.",
    "Secretaria do Tesouro Nacional [STN]. 2026. Capacidade de Pagamento de Municípios (CAPAG). Disponível em: "
    "<https://www.tesourotransparente.gov.br/ckan/dataset/capag-municipios>. Acesso em: 16 out. 2026.",
]


def clonar_paragrafo(modelo, texto, apos, alinhamento=None, negrito=None):
    """Insere após `apos` um parágrafo com a formatação de `modelo` e o texto dado."""
    novo = copy.deepcopy(modelo._p)
    for r in novo.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r'):
        novo.remove(r)
    apos._p.addnext(novo)
    from docx.text.paragraph import Paragraph
    p = Paragraph(novo, apos._parent)
    run = p.add_run(texto)
    run.font.name = "Arial"; run.font.size = Pt(11)
    if negrito is not None:
        run.bold = negrito
    if alinhamento is not None:
        p.alignment = alinhamento
    return p


def main():
    d = Document(TEMPLATE)
    P = d.paragraphs
    modelo_corpo = P[11]   # parágrafo de corpo justificado do template
    # Cabeçalho de identificação e título
    for idx, txt in [(0, f"Aluno(a): {ALUNO}"), (1, f"Orientador(a): {ORIENTADOR}"), (2, f"Curso: {CURSO}")]:
        p = P[idx]; p.runs[0].text = txt
        for r in p.runs[1:]: r.text = ""
    P[4].runs[0].text = TITULO
    for r in P[4].runs[1:]: r.text = ""

    secoes = {9: INTRODUCAO, 17: OBJETIVO, 22: METODOLOGIA, 31: RESULTADOS_ESPERADOS}
    # índices das instruções a remover (todos os parágrafos de corpo e avisos do template)
    remover = set([6, 7, 11, 12, 13, 14, 15, 19, 20, 24, 25, 26, 27, 28, 29, 33, 34, 38, 39, 44, 45])
    # inserir conteúdo logo após a linha em branco que segue cada título
    for idx_titulo, paras in secoes.items():
        ancora = P[idx_titulo + 1]
        for texto in paras:
            ancora = clonar_paragrafo(modelo_corpo, texto, ancora)
    # referências: alinhadas à esquerda, sem recuo, espaçamento simples, separadas por linha em branco
    ancora = P[43]
    for i, ref in enumerate(REFERENCIAS):
        ancora = clonar_paragrafo(modelo_corpo, ref, ancora, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
        pf = ancora.paragraph_format
        pf.first_line_indent = Cm(0); pf.left_indent = Cm(0); pf.line_spacing = 1.0
        pf.space_after = Pt(11); pf.space_before = Pt(0)
    # cronograma
    t = d.tables[0]
    def set_cell(cell, txt, center=True, size=10):
        par = cell.paragraphs[0]
        for r in par.runs: r.text = ""
        run = par.add_run(txt); run.font.name = "Arial"; run.font.size = Pt(size)
        if center: par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for j, m in enumerate(MESES, start=1):
        set_cell(t.rows[1].cells[j], m)
    for i, (atv, meses) in enumerate(ATIVIDADES, start=2):
        set_cell(t.rows[i].cells[0], atv, center=False)
        for j in range(1, 11):
            set_cell(t.rows[i].cells[j], "X" if j in meses else "")
    # remover instruções (de trás para frente para não deslocar índices)
    for idx in sorted(remover, reverse=True):
        el = P[idx]._p; el.getparent().remove(el)
    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    d.save(SAIDA)
    print("salvo em", SAIDA)
    print("palavras Introdução:", sum(len(p.split()) for p in INTRODUCAO),
          "| Metodologia:", sum(len(p.split()) for p in METODOLOGIA),
          "| Resultados esperados:", sum(len(p.split()) for p in RESULTADOS_ESPERADOS))


if __name__ == "__main__":
    main()
